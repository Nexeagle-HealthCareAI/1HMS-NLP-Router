import json
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from agentgate.config import load_dataset

ROOT = Path(__file__).resolve().parent.parent


def run_cli(tmp_path, *args, buggy=False):
    env = {**os.environ, "PYTHONPATH": str(ROOT / "src")}
    env["MOCK_AGENT_BUGGY"] = "1" if buggy else "0"
    return subprocess.run([sys.executable, "-m", "agentgate", *args, "--agent", "config/agent.demo.yaml",
                           "--out", str(tmp_path)], cwd=ROOT, env=env, capture_output=True, text=True)


def latest(tmp_path):
    return (json.loads((tmp_path / "latest" / "results.json").read_text()),
            tmp_path / "latest" / "report.pdf",
            tmp_path / "latest" / "report.xlsx")


def test_good_agent_smoke_passes_with_one_page(tmp_path):
    p = run_cli(tmp_path, "run", "--tier", "smoke")
    assert p.returncode == 0, p.stdout + p.stderr
    res, pdf, xlsx = latest(tmp_path)
    assert res["gate"]["verdict"] == "SMOKE PASS"
    assert len(PdfReader(str(pdf)).pages) == 1
    workbook = load_workbook(xlsx, data_only=True)
    assert workbook.sheetnames == ["Summary", "Golden Dataset", "Check Details"]
    cases = workbook["Golden Dataset"]
    assert [cell.value for cell in cases[1]] == [
        "Input", "Expected output", "Generated response", "Match or error",
        "% accuracy of response", "Pass / fail",
    ]
    golden_cases = load_dataset(str(ROOT / "dataset" / "golden_cases.json"))
    assert cases.max_row == len(golden_cases) + 1
    assert cases["A2"].value == "I'd like a refund for order ORD-1001."
    assert cases["B2"].value and cases["C2"].value
    assert 0 <= cases["E2"].value <= 1
    not_run = [row for row in cases.iter_rows(min_row=2, values_only=True) if row[5] == "NOT RUN"]
    assert not_run
    assert all(row[3] == "Not run for this tier" and row[4] is None for row in not_run)
    assert workbook["Check Details"].max_row > 1
    assert res["cases"][0]["expected"]
    assert res["cases"][0]["produced"]["text"]
    assert "actual" in res["cases"][0]["check_results"][0]
    assert "match_percentage" in res["gate"]["totals"]


def test_source_discovery_runs_python_entrypoint(tmp_path):
    env = {**os.environ, "PYTHONPATH": str(ROOT / "src")}
    result = subprocess.run(
        [sys.executable, "-m", "agentgate", "run", "--tier", "smoke",
         "--source", "demo/mock_agent.py", "--out", str(tmp_path)],
        cwd=ROOT, env=env, capture_output=True, text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    res, _, xlsx = latest(tmp_path)
    assert res["gate"]["verdict"] == "SMOKE PASS"
    assert "callable run" in res["meta"]["agent"]
    assert xlsx.exists()


def test_good_agent_full_without_judge_is_pending_manual(tmp_path):
    p = run_cli(tmp_path, "run", "--tier", "full", "--no-judge")
    assert p.returncode == 0, p.stdout + p.stderr
    res, _, _ = latest(tmp_path)
    assert res["gate"]["verdict"] == "PENDING MANUAL CHECKS"
    p = run_cli(tmp_path, "run", "--tier", "full", "--no-judge", "--rollback-rehearsed")
    assert latest(tmp_path)[0]["gate"]["verdict"] == "GO"


def test_buggy_agent_is_no_go_and_never_prints_the_secret(tmp_path):
    p = run_cli(tmp_path, "run", "--tier", "full", "--no-judge", buggy=True)
    assert p.returncode == 1
    res, pdf, _ = latest(tmp_path)
    assert res["gate"]["verdict"] == "NO-GO"
    assert res["gate"]["totals"]["must_never_events"] >= 1
    assert len(PdfReader(str(pdf)).pages) == 1
    assert "CANARY-7f3a" not in json.dumps(res) and "CANARY-7f3a" not in PdfReader(str(pdf)).pages[0].extract_text()


def test_unreachable_judge_fails_closed(tmp_path):
    env = {"JUDGE_BASE_URL": "http://127.0.0.1:9/v1"}
    os.environ.update(env)
    try:
        p = run_cli(tmp_path, "run", "--tier", "full")
    finally:
        os.environ.pop("JUDGE_BASE_URL")
    assert p.returncode == 1 and "NO-GO" in p.stdout


def test_dataset_validates():
    env = {**os.environ, "PYTHONPATH": str(ROOT / "src")}
    p = subprocess.run([sys.executable, "-m", "agentgate", "validate"], cwd=ROOT, env=env, capture_output=True, text=True)
    assert p.returncode == 0, p.stdout

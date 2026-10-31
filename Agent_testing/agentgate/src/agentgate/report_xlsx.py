"""Detailed Excel report with run summary, case outcomes, and assertion comparisons."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_PASS_FILL = PatternFill("solid", fgColor="E2F0D9")
_FAIL_FILL = PatternFill("solid", fgColor="FCE4D6")
_WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
_WHITE_BOLD = Font(color="FFFFFF", bold=True)
_MAX_CELL_LENGTH = 32000


def _value(value) -> str:
    if isinstance(value, str):
        rendered = value
    else:
        rendered = json.dumps(value, ensure_ascii=False, indent=2, default=str)
    if len(rendered) > _MAX_CELL_LENGTH:
        rendered = rendered[:_MAX_CELL_LENGTH] + "\n[truncated for Excel cell limit]"
    if rendered.startswith(("=", "+", "-", "@")):
        rendered = "'" + rendered
    return rendered


def _style_sheet(sheet, widths: list[int], auto_filter: bool = True) -> None:
    sheet.freeze_panes = "A2"
    if auto_filter:
        sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _WHITE_BOLD
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_view.showGridLines = False


def build_xlsx(res: dict, path: str) -> None:
    meta, gate = res["meta"], res["gate"]
    cases = res.get("report_cases", res["cases"])
    totals = gate["totals"]
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    summary.append(["AgentGate Evaluation Report", ""])
    summary.append(["Run ID", meta.get("run_id", "")])
    summary.append(["Tier", meta.get("tier", "")])
    summary.append(["Agent", meta.get("agent", "")])
    summary.append(["Dataset hash", meta.get("dataset_hash", "")])
    summary.append(["Thresholds hash", meta.get("thresholds_hash", "")])
    summary.append(["Commit", meta.get("commit", "")])
    summary.append(["AgentGate version", meta.get("version", "")])
    summary.append(["Verdict", gate.get("verdict", "")])
    summary.append(["Cases passed", f"{totals.get('passed', 0)} / {totals.get('cases', 0)}"])
    case_rate = totals.get("passed", 0) / totals["cases"] if totals.get("cases") else 0
    summary.append(["Tested case pass rate", case_rate])
    summary.append(["Assertion match", totals.get("match_percentage", 0) / 100])
    summary.append(["Checks passed", f"{totals.get('checks_passed', 0)} / {totals.get('checks_total', 0)}"])
    summary.append(["Agent/API error rate", totals.get("error_rate", 0)])
    summary.append(["p95 latency (ms)", totals.get("p95_latency_ms", 0)])
    summary.append(["Must-never events", totals.get("must_never_events", 0)])
    summary.append(["Judge average", totals.get("judge_avg")])
    summary.append(["Golden dataset entries", len(cases)])
    summary.append(["Cases not run", sum(1 for case in cases if case.get("report_status") == "NOT RUN")])
    summary.append([])
    category_header_row = summary.max_row + 1
    summary.append(["Category", "Cases", "Passed", "Match rate", "Required", "Status"])
    for category in gate.get("categories", []):
        summary.append([
            category["category"], category["cases"], category["passed"],
            category["rate"], category["required"], "PASS" if category["ok"] else "FAIL",
        ])
    summary.append([])
    summary.append(["Gate check", "Status", "Detail"])
    for check in gate.get("checks", []):
        summary.append([check["name"], check["status"].upper(), check["detail"]])

    summary["A1"].font = Font(size=16, bold=True, color="1F1F1F")
    for row in summary.iter_rows(min_row=2):
        if row[0].value in {"Tested case pass rate", "Assertion match", "Agent/API error rate"}:
            row[1].number_format = "0.0%"
    for row in summary.iter_rows(min_row=category_header_row, max_row=category_header_row):
        for cell in row:
            cell.fill = _HEADER_FILL
            cell.font = _WHITE_BOLD
    _style_sheet(summary, [34, 24, 55, 16, 16, 14], auto_filter=False)
    for row in summary.iter_rows(min_row=2):
        if row[0].value == "Verdict":
            row[1].fill = _PASS_FILL if "PASS" in str(row[1].value) or row[1].value == "GO" else _FAIL_FILL
        if row[0].value == "Status":
            row[1].fill = _PASS_FILL if row[1].value == "PASS" else _FAIL_FILL
        if row[1].value in {"WARN", "PENDING"}:
            row[1].fill = _WARN_FILL

    case_sheet = workbook.create_sheet("Golden Dataset")
    case_sheet.append([
        "Input", "Expected output", "Generated response", "Match or error",
        "% accuracy of response", "Pass / fail",
    ])
    for case in cases:
        produced = case.get("produced") or {}
        status = case.get("report_status")
        if status is None:
            status = "PASS" if case.get("passed") else "FAIL"
        checks = [check for check in case.get("check_results", []) if check.get("passed") < check.get("total")]
        if status == "NOT RUN":
            match_or_error = "Not run for this tier"
            accuracy = None
        elif case.get("error") or produced.get("error"):
            match_or_error = case.get("error") or produced["error"]
            accuracy = case.get("match_percentage", 0) / 100
        elif status == "PASS":
            match_or_error = "MATCH"
            accuracy = case.get("match_percentage", 0) / 100
        else:
            mismatch = "; ".join(
                f"{check.get('name')}: expected {_value(check.get('expected'))}, "
                f"got {_value(check.get('actual'))}"
                for check in checks
            )
            match_or_error = f"MISMATCH: {mismatch}" if mismatch else "Mismatch; see Check Details"
            accuracy = case.get("match_percentage", 0) / 100

        expected = {"expect": case.get("expected", {})}
        if case.get("source_constraints"):
            expected["constraints"] = case["source_constraints"]
        generated = produced.get("text") or ""
        if produced.get("raw") is not None:
            generated = _value(produced["raw"])
        elif produced.get("tool_calls") or produced.get("final_state"):
            generated = _value({
                "text": produced.get("text", ""),
                "tool_calls": produced.get("tool_calls", []),
                "final_state": produced.get("final_state", {}),
                "handoff": produced.get("handoff", False),
            })
        if produced.get("error"):
            generated = "No response"

        case_sheet.append([
            _value(case.get("input", "")),
            _value(expected),
            _value(generated),
            _value(match_or_error),
            accuracy,
            status,
        ])
    _style_sheet(case_sheet, [52, 58, 64, 64, 20, 16])
    for row in case_sheet.iter_rows(min_row=2):
        row[4].number_format = "0.0%"
        row[5].fill = (
            _PASS_FILL if row[5].value == "PASS"
            else _WARN_FILL if row[5].value == "NOT RUN"
            else _FAIL_FILL
        )

    detail_sheet = workbook.create_sheet("Check Details")
    detail_sheet.append([
        "Case ID", "Category", "Check", "Expected", "Produced / observed",
        "Match %", "Passed repeats", "Total repeats", "Result",
    ])
    for case in cases:
        for check in case.get("check_results", []):
            detail_sheet.append([
                case.get("id"), case.get("category"), check.get("name"),
                _value(check.get("expected")), _value(check.get("actual")),
                check.get("match_percentage", 0) / 100,
                check.get("passed", 0), check.get("total", 0),
                "PASS" if check.get("passed") == check.get("total") else "FAIL",
            ])
    _style_sheet(detail_sheet, [14, 18, 32, 42, 52, 12, 16, 16, 12])
    for row in detail_sheet.iter_rows(min_row=2):
        row[5].number_format = "0.0%"
        row[8].fill = _PASS_FILL if row[8].value == "PASS" else _FAIL_FILL

    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
